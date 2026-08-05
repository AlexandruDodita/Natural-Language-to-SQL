"""FastAPI layer: the SSE contract the frontend depends on must not change.

``chat-app`` parses ``data: [META]{...}``, ``data: [DATA]{...}``, JSON-encoded
text chunks, ``data: [DONE]`` and ``data: [ERROR] ...`` in that order, so these
tests pin the framing and the original ``sql_meta`` keys.
"""

from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

import main
from db import QueryResult


class ScriptedLLM:
    name = "scripted"

    def __init__(self, payload: dict, text: str = "Here are the results."):
        self.payload = payload
        self.text = text

    def generate_json(self, prompt: str) -> dict:
        return self.payload

    def stream_text(self, prompt, history=None):
        for word in self.text.split():
            yield word + " "


class FakeRunner:
    name = "fake"
    supports_session_context = False

    def __init__(self):
        self.executed: list[str] = []

    async def execute(self, sql: str, session=None):
        self.executed.append(sql)
        return QueryResult(
            columns=["city", "n"],
            rows=[["Miami", 3], ["Chicago", 5]],
            row_count=2,
            duration_ms=1.5,
        )


@pytest.fixture()
def client(catalog, index, policy_engine, monkeypatch, tmp_path):
    async def noop():
        return None

    monkeypatch.setattr(main.state, "initialize", noop)
    main.state.catalog = catalog
    main.state.index = index
    main.state.policy = policy_engine
    main.state.runner = FakeRunner()
    main.state.direct_runner = None
    main.state.llm = ScriptedLLM({"sql": "SELECT city FROM locations", "chart": None})
    main.state.schema_text_fallback = catalog.render()
    main.state.status = {"tables": len(catalog.tables)}
    main.settings.policy_enabled = True
    main.settings.dry_run_enabled = False
    main.telemetry.path = str(tmp_path / "telemetry.jsonl")
    main.telemetry.enabled = True
    with TestClient(main.app) as c:
        yield c


def sse_lines(text: str) -> list[str]:
    return [line[6:] for line in text.splitlines() if line.startswith("data: ")]


# ---------------------------------------------------------------------------
# Health / introspection endpoints
# ---------------------------------------------------------------------------
def test_health(client):
    body = client.get("/health").json()
    assert body["status"] == "ok"
    assert "model" in body


def test_config_exposes_the_ablation_knobs(client):
    knobs = client.get("/config").json()["knobs"]
    assert {"retrieval_top_k", "max_repair_attempts", "retrieval_fallback_max_tables"} <= set(knobs)


def test_schema_endpoint(client):
    body = client.get("/schema").json()
    assert body["database"] == "car_rental"
    assert len(body["tables"]) == 9


def test_retrieve_endpoint_reports_linked_tables(client):
    body = client.post("/retrieve", json={"question": "average salary per branch", "top_k": 2}).json()
    assert body["ranking"][0]["table"] == "employees"
    assert "tables" in body and body["prompt_text"]


def test_validate_endpoint_previews_the_policy_rewrite(client):
    body = client.post(
        "/validate",
        json={"sql": "SELECT id FROM vehicles", "user": {"role": "agent", "location_id": 4}},
    ).json()
    assert body["validation"]["ok"]
    assert "location_id = 4" in body["sql"]


def test_validate_endpoint_rejects_dangerous_sql(client):
    body = client.post("/validate", json={"sql": "DROP TABLE clients"}).json()
    assert body["validation"]["ok"] is False


# ---------------------------------------------------------------------------
# The SSE contract
# ---------------------------------------------------------------------------
def test_chat_sse_frames_are_unchanged(client):
    response = client.post(
        "/chat", json={"messages": [{"role": "user", "content": "cities?"}]}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")

    lines = sse_lines(response.text)
    assert lines[0].startswith("[META]")
    assert any(line.startswith("[DATA]") for line in lines)
    assert lines[-1] == "[DONE]"

    meta = json.loads(lines[0][6:])
    # the keys the frontend has always read
    assert set(meta) >= {"sql", "row_count", "duration_ms", "blocked"}
    assert meta["row_count"] == 2
    assert meta["blocked"] is None

    data = json.loads(next(line for line in lines if line.startswith("[DATA]"))[6:])
    assert data["columns"] == ["city", "n"]
    assert len(data["rows"]) == 2

    text_chunks = [
        json.loads(line) for line in lines[1:-1] if not line.startswith("[")
    ]
    assert "".join(text_chunks).strip() == "Here are the results."


def test_chat_meta_carries_the_new_observability_fields(client):
    response = client.post(
        "/chat", json={"messages": [{"role": "user", "content": "cities?"}]}
    )
    meta = json.loads(sse_lines(response.text)[0][6:])
    assert meta["outcome"] == "answered"
    assert meta["retrieval"]["tables"]
    assert meta["attempts"] == 1
    assert meta["role"] == "manager"
    assert meta["request_id"]


def test_chat_works_without_a_user_field(client):
    """The existing frontend sends only `messages`."""
    response = client.post(
        "/chat", json={"messages": [{"role": "user", "content": "cities?"}]}
    )
    assert response.status_code == 200
    assert "[DONE]" in response.text


def test_chat_applies_the_user_context_when_provided(client):
    response = client.post(
        "/chat",
        json={
            "messages": [{"role": "user", "content": "cars?"}],
            "user": {"user_id": "agent-7", "role": "agent", "location_id": 4},
        },
    )
    meta = json.loads(sse_lines(response.text)[0][6:])
    assert meta["role"] == "agent"


def test_chat_reports_a_policy_block(client, monkeypatch):
    main.state.llm = ScriptedLLM({"sql": "SELECT salary FROM employees", "chart": None})
    response = client.post(
        "/chat",
        json={
            "messages": [{"role": "user", "content": "salaries?"}],
            "user": {"role": "agent", "location_id": 4},
        },
    )
    lines = sse_lines(response.text)
    meta = json.loads(lines[0][6:])
    assert meta["outcome"] == "blocked_by_policy"
    assert "salary" in meta["blocked"]
    assert not any(line.startswith("[DATA]") for line in lines)
    assert lines[-1] == "[DONE]"


def test_chat_reports_a_clarification(client):
    main.state.llm = ScriptedLLM(
        {"clarification": "By total spend or by number of rentals?", "sql": None}
    )
    response = client.post(
        "/chat", json={"messages": [{"role": "user", "content": "best clients?"}]}
    )
    meta = json.loads(sse_lines(response.text)[0][6:])
    assert meta["outcome"] == "clarification"
    assert "total spend" in meta["clarification"]


def test_chat_rejects_an_empty_message_list(client):
    assert client.post("/chat", json={"messages": []}).status_code == 400


def test_chat_writes_telemetry(client, tmp_path):
    client.post("/chat", json={"messages": [{"role": "user", "content": "cities?"}]})
    records = main.telemetry.read_all()
    assert records
    assert records[-1]["outcome"] == "answered"
    assert "answer" in records[-1]["stages"]


def test_telemetry_endpoints(client):
    client.post("/chat", json={"messages": [{"role": "user", "content": "cities?"}]})
    assert client.get("/telemetry/recent").json()["records"]
    assert client.get("/telemetry/summary").json()["requests"] >= 1


# ---------------------------------------------------------------------------
# The Excel contract is untouched
# ---------------------------------------------------------------------------
def test_report_endpoint_still_produces_a_workbook(client):
    response = client.post(
        "/report",
        json={
            "columns": ["city", "n"],
            "rows": [["Miami", 3], ["Chicago", 5]],
            "chart": {"type": "bar", "title": "Per city", "x": "city", "y": "n"},
            "title": "Report",
        },
    )
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Data", "Chart"]
    assert [c.value for c in wb["Data"][1]] == ["city", "n"]


def test_report_without_a_chart_has_one_sheet(client):
    response = client.post(
        "/report", json={"columns": ["a"], "rows": [[1]], "title": "x"}
    )
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Data"]
