#!/usr/bin/env python3
"""Scores the automatic Excel report: chart choice, and workbook structure.

    GEMINI_API_KEY=... python benchmark/report_score.py --model gemini-2.5-flash
    python benchmark/report_score.py --local --local-model qwen3.5-9b

The thesis claims the system exports a result set as a two-sheet Excel workbook
with an appropriate chart. That claim has two independent halves and they fail
in different ways, so they are measured separately:

chart-type accuracy
    a judgement call the model makes. Scored against ``expected_chart`` in
    questions.yaml, which lists every defensible type for that result shape - a
    five-row breakdown is as legitimate a pie as it is a bar, and marking one
    of them wrong would measure taste rather than capability.

structural validity
    a property of the generated file, not of the model's taste. The workbook is
    built with the product's own ``rag-service/report.py`` and then read back
    with openpyxl: both sheets present, the Data sheet equal to the result set
    the query returned, a chart object of the requested class actually embedded,
    and its series and category references pointing at the columns the model
    named. A workbook that opens but charts the wrong column is a silent
    failure in the product, and only a read-back catches it.

    Read this figure next to the chart-type column, never on its own. It asks
    whether the file matches what the model asked for, so a model that asks for
    no chart at all passes it trivially.

Chart selection is scored on gold SQL, deliberately. The question is whether a
model can choose a visualisation for a result set, and running it on the
model's own SQL would score a model that wrote the wrong query twice.
"""

from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import sys

import yaml

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "rag-service"))

import clients  # noqa: E402

MAX_ROWS = 200  # the product's own row cap for a report
SAMPLE_ROWS = 8  # rows shown to the model; it needs the shape, not the data


def build_chart_prompt(question: str, sql: str, columns: list[str], rows: list[list]) -> str:
    """The product's chart rules, applied to a result set that already exists.

    CHART_RULES is imported rather than restated so that this scorer can never
    drift from the prompt the running system uses.
    """
    from llm import CHART_RULES  # rag-service/llm.py

    sample = [[str(v) for v in r] for r in rows[:SAMPLE_ROWS]]
    return "\n\n".join([
        "You choose the visualisation for a database query result that will be "
        "exported to an Excel report.",
        f"User question: {question}",
        f"SQL that produced the result:\n{sql}",
        f"Result columns: {json.dumps(columns)}",
        f"First rows: {json.dumps(sample, ensure_ascii=False)}",
        f"Total rows returned: {len(rows)}",
        CHART_RULES,
        'Respond ONLY with a JSON object, no markdown and no explanation:\n'
        '{"type": "bar|line|pie|area|none", "title": "...", "x": "column_name", "y": "column_name"}',
    ])


def parse_chart_json(text: str) -> dict | None:
    t = (text or "").strip()
    if t.startswith("```"):
        t = re.sub(r"^```[a-zA-Z]*\n?", "", t)
        t = re.sub(r"\n?```$", "", t).strip()
    m = re.search(r"\{.*\}", t, re.S)
    if not m:
        return None
    try:
        obj = json.loads(m.group(0))
    except Exception:
        return None
    return obj if isinstance(obj, dict) else None


# ---------------------------------------------------------------------------
# Structural validation of the produced workbook
# ---------------------------------------------------------------------------
CHART_CLASSES = {"bar": "BarChart", "line": "LineChart", "pie": "PieChart", "area": "AreaChart"}


def _cell_eq(written, original) -> bool:
    """Excel round-trips a value through its own type system; compare loosely.

    A workbook cell is not required to hold the same Python object the driver
    returned - openpyxl writes a ``Decimal`` as a float and reads a ``date``
    back as a ``datetime``. Reusing the execution scorer's normalisation keeps
    "the Data sheet matches the query result" meaning the same thing here as it
    does there, instead of failing a correct workbook on a type round-trip.
    """
    from score import _norm

    a, b = _norm(written), _norm(original)
    if a == b:
        return True
    # A date written to Excel comes back as midnight on that date.
    if isinstance(a, str) and isinstance(b, str):
        return a.replace(" 00:00:00", "") == b.replace(" 00:00:00", "")
    return False


def validate_workbook(xlsx: bytes, columns: list[str], rows: list[list],
                      chart: dict | None) -> dict:
    """Read the workbook back and check every claim the thesis makes about it."""
    import io

    import openpyxl

    checks: dict[str, object] = {"opens": False, "data_sheet": False, "header_ok": False,
                                 "data_ok": False, "chart_sheet_ok": False,
                                 "chart_embedded": False, "chart_type_ok": False,
                                 "chart_refs_ok": False, "notes": []}
    notes: list[str] = checks["notes"]  # type: ignore[assignment]
    want_chart = bool(chart) and chart.get("type") in CHART_CLASSES
    # build_excel only draws a chart when both named columns exist; a model that
    # invents a column name gets a chartless workbook, which is a failure of the
    # model and not of the writer, and is recorded as such.
    axes_resolvable = want_chart and chart.get("x") in columns and chart.get("y") in columns
    checks["axes_resolvable"] = axes_resolvable

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    except Exception as e:
        notes.append(f"load failed: {type(e).__name__}: {e}")
        return checks
    checks["opens"] = True

    if "Data" not in wb.sheetnames:
        notes.append("no Data sheet")
        return checks
    checks["data_sheet"] = True
    ws = wb["Data"]

    header = [c.value for c in ws[1]]
    checks["header_ok"] = header == columns
    if not checks["header_ok"]:
        notes.append(f"header {header} != {columns}")

    body = list(ws.iter_rows(min_row=2, values_only=True))
    if len(body) != len(rows):
        notes.append(f"{len(body)} data rows, expected {len(rows)}")
    else:
        checks["data_ok"] = all(
            len(w) == len(o) and all(_cell_eq(a, b) for a, b in zip(w, o))
            for w, o in zip(body, rows)
        )
        if not checks["data_ok"]:
            notes.append("cell values differ from the query result")

    has_chart_sheet = "Chart" in wb.sheetnames
    checks["chart_sheet_ok"] = has_chart_sheet == axes_resolvable
    if not checks["chart_sheet_ok"]:
        notes.append("Chart sheet present" if has_chart_sheet else "Chart sheet missing")

    if not axes_resolvable:
        if want_chart:
            notes.append(f"axes {chart.get('x')!r}/{chart.get('y')!r} not among {columns}")
        # No chart wanted (or none drawable): the chart checks do not apply.
        checks["chart_embedded"] = checks["chart_type_ok"] = checks["chart_refs_ok"] = not want_chart
        return checks

    objs = list(getattr(wb["Chart"], "_charts", []))
    checks["chart_embedded"] = len(objs) == 1
    if not objs:
        notes.append("Chart sheet holds no chart object")
        return checks

    obj = objs[0]
    checks["chart_type_ok"] = type(obj).__name__ == CHART_CLASSES[chart["type"]]
    if not checks["chart_type_ok"]:
        notes.append(f"{type(obj).__name__} embedded, expected {CHART_CLASSES[chart['type']]}")

    # The references must point at the columns the model named, over exactly the
    # rows the query returned.
    from openpyxl.utils import get_column_letter

    x_col = get_column_letter(columns.index(chart["x"]) + 1)
    y_col = get_column_letter(columns.index(chart["y"]) + 1)
    last = len(rows) + 1
    want_val = f"'Data'!${y_col}$2:${y_col}${last}"
    want_cat = f"'Data'!${x_col}$2:${x_col}${last}"
    try:
        ser = obj.series[0]
        got_val = str(ser.val.numRef.f)
        got_cat = str((ser.cat.numRef or ser.cat.strRef).f)
    except Exception as e:
        notes.append(f"series unreadable: {type(e).__name__}: {e}")
        return checks
    def collapse(ref: str) -> str:
        # openpyxl writes a one-row range as a single cell ($B$2, not $B$2:$B$2).
        # That is the same reference, so comparing the raw strings would fail a
        # correct chart drawn over a single-row result.
        left, _, right = ref.partition(":")
        return left if right and left.endswith(right) else ref

    checks["chart_refs_ok"] = (collapse(got_val) == collapse(want_val)
                               and collapse(got_cat) == collapse(want_cat))
    if not checks["chart_refs_ok"]:
        notes.append(f"refs val={got_val} cat={got_cat}, expected val={want_val} cat={want_cat}")
    return checks


def summarise_report(records: list[dict]) -> dict:
    n = len(records)
    n_type = sum(1 for r in records if r["chart_type_correct"])
    n_valid = sum(1 for r in records if r["structurally_valid"])
    chartable = [r for r in records if r["expected_chart"] != ["none"]]
    scalar = [r for r in records if r["expected_chart"] == ["none"]]
    import clients as _clients
    return {
        "n": n,
        "chart_type_correct": n_type,
        "chart_type_accuracy": round(n_type / n, 4) if n else 0.0,
        "chart_type_accuracy_chartable": round(
            sum(1 for r in chartable if r["chart_type_correct"]) / len(chartable), 4) if chartable else None,
        "chart_type_accuracy_scalar": round(
            sum(1 for r in scalar if r["chart_type_correct"]) / len(scalar), 4) if scalar else None,
        "structurally_valid": n_valid,
        "structural_validity": round(n_valid / n, 4) if n else 0.0,
        "unparseable_responses": sum(1 for r in records if r["predicted_chart"] is None),
        "usage": _clients.aggregate_usage([r["usage"] for r in records if r.get("usage")]),
    }


def revalidate(path: pathlib.Path, questions_path: pathlib.Path,
               database_url: str | None) -> int:
    """Redo the workbook checks on a finished run without calling any model.

    The chart the model chose is already recorded, and building and reading back
    a workbook is deterministic, so a fix to the checker does not require paying
    for the model's judgement a second time.
    """
    import psycopg

    from report import ChartConfig, ReportRequest, build_excel
    from run import ensure_database

    doc = json.loads(path.read_text())
    gold = {q["id"]: q for q in yaml.safe_load(questions_path.read_text())["questions"]}
    uri = ensure_database(database_url, REPO)

    with psycopg.connect(uri, autocommit=True) as conn:
        conn.execute("SET default_transaction_read_only = on")
        for rec in doc["results"]:
            q = gold[rec["id"]]
            with conn.cursor() as cur:
                cur.execute(q["gold_sql"])
                columns = [d.name for d in cur.description]
                rows = [list(r) for r in cur.fetchmany(MAX_ROWS)]
            chart = rec.get("raw_chart") or {}
            predicted = rec.get("predicted_chart")
            cfg = None
            if predicted in CHART_CLASSES:
                cfg = ChartConfig(type=predicted,
                                  title=str(chart.get("title") or q["question"])[:80],
                                  x=str(chart.get("x") or ""), y=str(chart.get("y") or ""))
            try:
                xlsx = build_excel(ReportRequest(columns=columns, rows=rows, chart=cfg,
                                                 title=q["id"]))
                checks = validate_workbook(xlsx, columns, rows,
                                           cfg.model_dump() if cfg else None)
                build_err = None
            except Exception as e:
                checks = {"opens": False, "notes": [f"{type(e).__name__}: {e}"]}
                build_err = f"{type(e).__name__}: {e}"
            rec["checks"] = checks
            rec["structurally_valid"] = structurally_valid(checks)
            if build_err:
                rec["error"] = build_err
            print(f"  {rec['id']} {'valid' if rec['structurally_valid'] else 'INVALID'} "
                  f"{'; '.join(checks.get('notes') or [])[:90]}")

    doc["summary"] = summarise_report(doc["results"])
    path.write_text(json.dumps(doc, indent=2, default=str))
    s = doc["summary"]
    print(f"\n{path.name}: chart-type {s['chart_type_accuracy']:.1%}, "
          f"structural validity {s['structural_validity']:.1%} "
          f"({s['structurally_valid']}/{s['n']})")
    return 0


def structurally_valid(checks: dict) -> bool:
    keys = ["opens", "data_sheet", "header_ok", "data_ok", "chart_sheet_ok",
            "chart_embedded", "chart_type_ok", "chart_refs_ok"]
    return all(bool(checks.get(k)) for k in keys)


# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--questions", type=pathlib.Path, default=HERE / "questions.yaml")
    ap.add_argument("--database-url", default=os.environ.get("BENCHMARK_DATABASE_URL"))
    ap.add_argument("--model", help="Gemini model id (default: $GEMINI_MODEL)")
    ap.add_argument("--local", action="store_true", help="use the OpenAI-compatible local endpoint")
    ap.add_argument("--local-base-url", default=clients.DEFAULT_LOCAL_BASE_URL)
    ap.add_argument("--local-model", default=clients.DEFAULT_LOCAL_MODEL)
    ap.add_argument("--limit", type=int)
    ap.add_argument("--out", type=pathlib.Path)
    ap.add_argument("--revalidate", type=pathlib.Path,
                    help="re-run only the workbook checks on an existing results "
                         "file, reusing the chart the model already chose; no "
                         "model is called")
    args = ap.parse_args()

    from report import ChartConfig, ReportRequest, build_excel  # rag-service/report.py

    import psycopg

    from run import ensure_database  # reuses the same fixture the SQL benchmark uses

    if args.revalidate:
        return revalidate(args.revalidate, args.questions, args.database_url)

    questions = [q for q in yaml.safe_load(args.questions.read_text())["questions"]
                 if q.get("gold_sql") and q.get("expected_chart")]
    if args.limit:
        questions = questions[: args.limit]

    label = args.local_model if args.local else (args.model or clients.gemini_model_name())
    uri = ensure_database(args.database_url, REPO)

    records: list[dict] = []
    with psycopg.connect(uri, autocommit=True) as conn:
        conn.execute("SET default_transaction_read_only = on")
        for i, q in enumerate(questions, 1):
            with conn.cursor() as cur:
                cur.execute(q["gold_sql"])
                columns = [d.name for d in cur.description]
                rows = [list(r) for r in cur.fetchmany(MAX_ROWS)]

            prompt = build_chart_prompt(q["question"], q["gold_sql"], columns, rows)
            try:
                if args.local:
                    text, usage = clients.local_complete(
                        prompt, base_url=args.local_base_url, model=args.local_model)
                else:
                    text, usage = clients.gemini_complete(prompt, model=args.model)
                err = None
            except Exception as e:
                text, usage, err = "", clients.Usage(), f"{type(e).__name__}: {e}"

            chart = parse_chart_json(text)
            predicted = (chart or {}).get("type") or ("none" if chart else None)
            expected = q["expected_chart"]
            type_ok = predicted in expected if predicted else False

            cfg = None
            if chart and predicted in CHART_CLASSES:
                cfg = ChartConfig(type=predicted, title=str(chart.get("title") or q["question"])[:80],
                                  x=str(chart.get("x") or ""), y=str(chart.get("y") or ""))
            try:
                xlsx = build_excel(ReportRequest(columns=columns, rows=rows, chart=cfg,
                                                 title=q["id"]))
                checks = validate_workbook(xlsx, columns, rows,
                                           cfg.model_dump() if cfg else None)
                build_err = None
            except Exception as e:
                checks, build_err = {"opens": False, "notes": [f"{type(e).__name__}: {e}"]}, str(e)

            records.append({
                "id": q["id"], "category": q["category"], "question": q["question"],
                "expected_chart": expected, "predicted_chart": predicted,
                "chart_type_correct": type_ok, "raw_chart": chart,
                "structurally_valid": structurally_valid(checks),
                "checks": checks, "error": err or build_err,
                "usage": usage.as_dict() if usage else None,
            })
            mark = "ok " if type_ok else "MISS"
            sv = "valid" if records[-1]["structurally_valid"] else "INVALID"
            print(f"  [{i:>2}/{len(questions)}] {q['id']} chart={str(predicted):<5} {mark} "
                  f"xlsx={sv:<7} exp={'/'.join(expected)}")

    summary = summarise_report(records)
    n, n_type, n_valid = summary["n"], summary["chart_type_correct"], summary["structurally_valid"]

    out = args.out or (HERE / "results" / f"report-{label}.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"model": label, "summary": summary, "results": records},
                              indent=2, default=str))

    print(f"\n=== excel report quality: {label} ===")
    print(f"chart-type accuracy   : {summary['chart_type_accuracy']:.1%} ({n_type}/{n})")
    print(f"  chartable results   : {summary['chart_type_accuracy_chartable']}")
    print(f"  scalar/text results : {summary['chart_type_accuracy_scalar']}")
    print(f"structural validity   : {summary['structural_validity']:.1%} ({n_valid}/{n})")
    print(f"unparseable responses : {summary['unparseable_responses']}")
    print(f"\nwrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
